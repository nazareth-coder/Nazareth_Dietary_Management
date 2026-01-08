import openpyxl
import bcrypt
from datetime import datetime
from typing import Optional, Dict, List, Tuple

from constants import EXCEL_FILE


USERS_SHEET_NAME = 'Users'


def _ensure_workbook() -> openpyxl.Workbook:
    """Load the main workbook, creating it if needed."""
    try:
        wb = openpyxl.load_workbook(EXCEL_FILE)
    except Exception:
        wb = openpyxl.Workbook()
    return wb


def ensure_users_sheet() -> None:
    """Ensure the Users sheet exists with a default admin user.

    Default admin credentials (you should change password after first login):
      username: admin
      password: admin123
    """
    wb = _ensure_workbook()

    # Create sheet and header if missing
    if USERS_SHEET_NAME not in wb.sheetnames:
        ws = wb.create_sheet(USERS_SHEET_NAME)
        ws.append(['username', 'password_hash', 'role', 'full_name', 'last_login'])

        # Create default admin user (admin / admin123)
        default_hash = bcrypt.hashpw('admin123'.encode('utf-8'), bcrypt.gensalt()).decode('utf-8')
        ws.append(['admin', default_hash, 'admin', 'System Administrator', ''])
        wb.save(EXCEL_FILE)
        return

    # If sheet exists but has no users, ensure at least one admin
    ws = wb[USERS_SHEET_NAME]
    has_any_user = False
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row and row[0]:
            has_any_user = True
            break
    if not has_any_user:
        default_hash = bcrypt.hashpw('admin123'.encode('utf-8'), bcrypt.gensalt()).decode('utf-8')
        ws.append(['admin', default_hash, 'admin', 'System Administrator', ''])
        wb.save(EXCEL_FILE)


def _verify_password(stored_hash: str, provided_password: str) -> bool:
    try:
        return bcrypt.checkpw(provided_password.encode('utf-8'), stored_hash.encode('utf-8'))
    except Exception:
        return False


def authenticate_user(username: str, password: str) -> Optional[Dict[str, str]]:
    """Authenticate against the Users sheet.

    Returns a dict with keys: username, role, full_name on success, or None on failure.
    """
    wb = None
    try:
        wb = openpyxl.load_workbook(EXCEL_FILE, read_only=True)
        if USERS_SHEET_NAME not in wb.sheetnames:
            return None
        ws = wb[USERS_SHEET_NAME]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row:
                continue
            stored_username, stored_hash, role, full_name, _ = (list(row) + [None] * 5)[:5]
            if stored_username == username and stored_hash and _verify_password(stored_hash, password):
                # Update last_login in a separate writeable session
                _update_last_login(username)
                return {
                    'username': stored_username,
                    'role': role or 'user',
                    'full_name': full_name or stored_username,
                }
        return None
    except Exception:
        return None
    finally:
        if wb is not None:
            try:
                wb.close()
            except Exception:
                pass


def _update_last_login(username: str) -> None:
    try:
        wb = openpyxl.load_workbook(EXCEL_FILE)
        if USERS_SHEET_NAME not in wb.sheetnames:
            return
        ws = wb[USERS_SHEET_NAME]
        for row in ws.iter_rows(min_row=2):
            if row[0].value == username:
                row[4].value = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
                break
        wb.save(EXCEL_FILE)
    except Exception:
        pass
    finally:
        try:
            wb.close()
        except Exception:
            pass


def create_user(username: str, password: str, role: str = 'user', full_name: str = '') -> bool:
    """Create a new user in the Users sheet. Returns False if username already exists."""
    try:
        wb = _ensure_workbook()
        if USERS_SHEET_NAME not in wb.sheetnames:
            ensure_users_sheet()
            wb = openpyxl.load_workbook(EXCEL_FILE)
        ws = wb[USERS_SHEET_NAME]

        # Ensure header exists
        if ws.max_row == 1 and ws['A1'].value != 'username':
            ws.append(['username', 'password_hash', 'role', 'full_name', 'last_login'])

        # Check for existing username
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row and row[0] == username:
                return False

        pw_hash = bcrypt.hashpw(password.encode('utf-8'), bcrypt.gensalt()).decode('utf-8')
        ws.append([username, pw_hash, role or 'user', full_name or username, ''])
        wb.save(EXCEL_FILE)
        return True
    except Exception:
        return False
    finally:
        try:
            wb.close()
        except Exception:
            pass


def change_password(username: str, old_password: str, new_password: str) -> bool:
    """Change an existing user's password. Returns True on success."""
    try:
        wb = openpyxl.load_workbook(EXCEL_FILE)
        if USERS_SHEET_NAME not in wb.sheetnames:
            return False
        ws = wb[USERS_SHEET_NAME]
        for row in ws.iter_rows(min_row=2):
            if row[0].value == username and row[1].value:
                if not _verify_password(str(row[1].value), old_password):
                    return False
                new_hash = bcrypt.hashpw(new_password.encode('utf-8'), bcrypt.gensalt()).decode('utf-8')
                row[1].value = new_hash
                wb.save(EXCEL_FILE)
                return True
        return False
    except Exception:
        return False
    finally:
        try:
            wb.close()
        except Exception:
            pass


def list_users() -> List[Dict[str, str]]:
    try:
        wb = openpyxl.load_workbook(EXCEL_FILE, read_only=True)
        if USERS_SHEET_NAME not in wb.sheetnames:
            return []
        ws = wb[USERS_SHEET_NAME]

        users: List[Dict[str, str]] = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row:
                continue
            username, _, role, full_name, last_login = (list(row) + [None] * 5)[:5]
            if not username:
                continue
            users.append({
                'username': str(username),
                'role': (str(role) if role else 'user'),
                'full_name': (str(full_name) if full_name else str(username)),
                'last_login': (str(last_login) if last_login else ''),
            })
        return users
    except Exception:
        return []
    finally:
        try:
            wb.close()
        except Exception:
            pass


def _find_user_row(ws, username: str):
    for row in ws.iter_rows(min_row=2):
        if row and row[0].value == username:
            return row
    return None


def _count_admins(ws) -> int:
    count = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        username, _, role, _, _ = (list(row) + [None] * 5)[:5]
        if not username:
            continue
        if (str(role).strip().lower() if role else 'user') == 'admin':
            count += 1
    return count


def update_user(username: str, new_username: Optional[str] = None, new_role: Optional[str] = None, new_full_name: Optional[str] = None) -> Tuple[bool, str]:
    try:
        wb = openpyxl.load_workbook(EXCEL_FILE)
        if USERS_SHEET_NAME not in wb.sheetnames:
            return False, 'Users sheet not found.'
        ws = wb[USERS_SHEET_NAME]

        row = _find_user_row(ws, username)
        if row is None:
            return False, 'User not found.'

        target_username = (new_username.strip() if isinstance(new_username, str) else '')
        if target_username and target_username != username:
            for existing in ws.iter_rows(min_row=2, values_only=True):
                if existing and existing[0] == target_username:
                    return False, 'Username already exists.'
            row[0].value = target_username

        if isinstance(new_role, str) and new_role.strip() != '':
            row[2].value = new_role.strip().lower()

        if isinstance(new_full_name, str) and new_full_name.strip() != '':
            row[3].value = new_full_name.strip()

        wb.save(EXCEL_FILE)
        return True, ''
    except Exception as e:
        return False, str(e)
    finally:
        try:
            wb.close()
        except Exception:
            pass


def set_user_password(username: str, new_password: str) -> Tuple[bool, str]:
    try:
        if not new_password or len(new_password) < 6:
            return False, 'Password must be at least 6 characters long.'
        wb = openpyxl.load_workbook(EXCEL_FILE)
        if USERS_SHEET_NAME not in wb.sheetnames:
            return False, 'Users sheet not found.'
        ws = wb[USERS_SHEET_NAME]
        row = _find_user_row(ws, username)
        if row is None:
            return False, 'User not found.'
        new_hash = bcrypt.hashpw(new_password.encode('utf-8'), bcrypt.gensalt()).decode('utf-8')
        row[1].value = new_hash
        wb.save(EXCEL_FILE)
        return True, ''
    except Exception as e:
        return False, str(e)
    finally:
        try:
            wb.close()
        except Exception:
            pass


def delete_user(username: str, requesting_username: Optional[str] = None) -> Tuple[bool, str]:
    try:
        if requesting_username and username == requesting_username:
            return False, 'You cannot delete your own account while logged in.'

        wb = openpyxl.load_workbook(EXCEL_FILE)
        if USERS_SHEET_NAME not in wb.sheetnames:
            return False, 'Users sheet not found.'
        ws = wb[USERS_SHEET_NAME]

        row = _find_user_row(ws, username)
        if row is None:
            return False, 'User not found.'

        role_value = (str(row[2].value).strip().lower() if row[2].value else 'user')
        if role_value == 'admin' and _count_admins(ws) <= 1:
            return False, 'Cannot delete the last admin user.'

        ws.delete_rows(row[0].row, 1)
        wb.save(EXCEL_FILE)
        return True, ''
    except Exception as e:
        return False, str(e)
    finally:
        try:
            wb.close()
        except Exception:
            pass
